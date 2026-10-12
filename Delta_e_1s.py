"""Sample a fixed video ROI every 1 second and export CIE76 color change.

Install: pip install opencv-python numpy openpyxl
Run:     python deltaE_video_roi.py [path/to/video.mp4]

The preview starts paused at the first frame. Space plays/pauses; drag a
rectangle on any preview frame and press Enter to run the analysis. The chosen
frame only defines the ROI: measurements always begin at video time 0 s.
"""

import argparse
import math
from pathlib import Path
import time

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


DISPLAY_MAX_WIDTH = 1400
DISPLAY_MAX_HEIGHT = 800
VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v"}


def mean_rgb_of_roi(frame_bgr, roi):
    x, y, w, h = roi
    image_h, image_w = frame_bgr.shape[:2]
    if w <= 0 or h <= 0 or x < 0 or y < 0 or x + w > image_w or y + h > image_h:
        raise ValueError(f"ROI {roi} is outside frame size {image_w} x {image_h}.")
    # Mean of all ROI pixels in sRGB, matching the previous PNG program.
    return frame_bgr[y:y + h, x:x + w, :3].mean(axis=(0, 1))[::-1]


def rgb_to_lab(mean_rgb):
    """Convert mean 8-bit sRGB through D65 XYZ into CIELAB (2-degree observer)."""
    rgb = np.asarray(mean_rgb, dtype=np.float64) / 255.0
    linear = np.where(rgb <= 0.04045, rgb / 12.92,
                      ((rgb + 0.055) / 1.055) ** 2.4)
    # sRGB -> XYZ (D65), using the same coefficients as scikit-image.
    xyz = np.array([
        [0.412453, 0.357580, 0.180423],
        [0.212671, 0.715160, 0.072169],
        [0.019334, 0.119193, 0.950227],
    ]) @ linear
    xyz /= np.array([0.95047, 1.0, 1.08883])
    delta = 6.0 / 29.0
    f = np.where(xyz > delta ** 3, np.cbrt(xyz),
                 xyz / (3 * delta ** 2) + 4.0 / 29.0)
    return np.array([116 * f[1] - 16, 500 * (f[0] - f[1]),
                     200 * (f[1] - f[2])])


def delta_e_cie76(lab, reference_lab):
    return float(np.linalg.norm(lab - reference_lab))


def video_info(cap):
    fps = float(cap.get(cv2.CAP_PROP_FPS))
    frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    if not math.isfinite(fps) or fps <= 0 or frame_count < 0:
        raise ValueError("Cannot read video FPS/frame count; this file may be unsupported.")
    return fps, frame_count


class RoiPlayer:
    def __init__(self, cap, first_frame, fps, frame_count):
        self.cap = cap
        self.fps = fps
        self.frame_count = frame_count
        self.frame = first_frame
        self.frame_index = 0
        h, w = first_frame.shape[:2]
        self.scale = min(1.0, DISPLAY_MAX_WIDTH / w, DISPLAY_MAX_HEIGHT / h)
        self.display_w = max(1, round(w * self.scale))
        self.display_h = max(1, round(h * self.scale))
        self.paused = True
        self.drag_start = None
        self.drag_end = None
        self.selected = None
        self.window = "Video ROI | Space: play/pause | drag: ROI | Enter: analyze | Esc: cancel"

    def mouse(self, event, x, y, flags, userdata):
        x = max(0, min(x, self.display_w - 1))
        y = max(0, min(y, self.display_h - 1))
        if event == cv2.EVENT_LBUTTONDOWN:
            self.paused = True
            self.drag_start = (x, y)
            self.drag_end = (x, y)
            self.selected = None
        elif event == cv2.EVENT_MOUSEMOVE and self.drag_start is not None:
            self.drag_end = (x, y)
        elif event == cv2.EVENT_LBUTTONUP and self.drag_start is not None:
            self.drag_end = (x, y)
            x0, x1 = sorted((self.drag_start[0], x))
            y0, y1 = sorted((self.drag_start[1], y))
            self.selected = (x0, y0, x1, y1) if x1 > x0 and y1 > y0 else None
            self.drag_start = None

    def draw(self):
        if (self.display_w, self.display_h) == (self.frame.shape[1], self.frame.shape[0]):
            shown = self.frame.copy()
        else:
            shown = cv2.resize(self.frame, (self.display_w, self.display_h),
                               interpolation=cv2.INTER_AREA)
        corners = self.selected
        if self.drag_start is not None:
            x0, x1 = sorted((self.drag_start[0], self.drag_end[0]))
            y0, y1 = sorted((self.drag_start[1], self.drag_end[1]))
            corners = (x0, y0, x1, y1)
        if corners is not None:
            cv2.rectangle(shown, corners[:2], corners[2:], (0, 255, 0), 2)
        status = "PAUSED" if self.paused else "PLAYING"
        label = f"{status}  {self.frame_index / self.fps:.2f} s  "
        label += "| Enter: confirm ROI" if self.selected else "| Drag to select ROI"
        cv2.rectangle(shown, (0, 0), (min(self.display_w - 1, 610), 31), (0, 0, 0), -1)
        cv2.putText(shown, label, (8, 23), cv2.FONT_HERSHEY_SIMPLEX,
                    0.58, (255, 255, 255), 1, cv2.LINE_AA)
        cv2.imshow(self.window, shown)

    def select(self):
        cv2.namedWindow(self.window, cv2.WINDOW_AUTOSIZE)
        cv2.setMouseCallback(self.window, self.mouse)
        try:
            next_frame_at = time.monotonic()
            while True:
                if not self.paused and time.monotonic() >= next_frame_at:
                    success, frame = self.cap.read()
                    if success:
                        self.frame = frame
                        self.frame_index += 1
                        next_frame_at = time.monotonic() + 1.0 / self.fps
                    else:
                        self.paused = True
                self.draw()
                key = cv2.waitKeyEx(15)
                if key in (27,):
                    raise RuntimeError("ROI selection cancelled.")
                if key in (13, 10) and self.selected is not None:
                    break
                if key in (32, ord('p'), ord('P')):
                    self.paused = not self.paused
                    next_frame_at = time.monotonic() + 1.0 / self.fps
                # Closing the preview also cancels instead of looping forever.
                if cv2.getWindowProperty(self.window, cv2.WND_PROP_VISIBLE) < 1:
                    raise RuntimeError("Preview window closed before selecting an ROI.")
        finally:
            cv2.destroyWindow(self.window)

        x0, y0, x1, y1 = self.selected
        image_h, image_w = self.frame.shape[:2]
        # Scale both corners separately to minimize rounding drift.
        left = max(0, min(image_w - 1, round(x0 / self.scale)))
        top = max(0, min(image_h - 1, round(y0 / self.scale)))
        right = max(left + 1, min(image_w, round((x1 + 1) / self.scale)))
        bottom = max(top + 1, min(image_h, round((y1 + 1) / self.scale)))
        return (left, top, right - left, bottom - top)


def analyze_video(cap, roi, interval_s=1.0):
    """Read frames nearest to 0, interval, 2*interval... on the FPS timeline."""
    if interval_s <= 0 or not math.isfinite(interval_s):
        raise ValueError("Sampling interval must be a positive finite number.")
    fps, frame_count = video_info(cap)
    rows = []
    reference_lab = None
    sample_number = 0
    last_frame_index = -1
    while True:
        target_time = sample_number * interval_s
        target_index = round(target_time * fps)
        if target_index <= last_frame_index:
            raise ValueError("Sampling interval is shorter than one video frame.")
        if frame_count > 0 and target_index >= frame_count:
            break
        if not cap.set(cv2.CAP_PROP_POS_FRAMES, target_index):
            raise RuntimeError(f"Cannot seek to {target_time:g} s (frame {target_index}).")
        success, frame = cap.read()
        if not success:
            if sample_number == 0:
                raise RuntimeError("Cannot decode the video at time 0 s.")
            if frame_count > 0:
                print(f"Warning: video ended or decode failed near {target_time:g} s.")
            break
        # POS_FRAMES is the next frame after read. Detect a decoder that did
        # not seek accurately rather than silently labeling the wrong frame.
        next_pos = float(cap.get(cv2.CAP_PROP_POS_FRAMES))
        actual_index = int(round(next_pos)) - 1 if math.isfinite(next_pos) and next_pos >= 1 else target_index
        if abs(actual_index - target_index) > 1:
            raise RuntimeError(f"Decoder sought to frame {actual_index}, expected {target_index}.")
        mean_rgb = mean_rgb_of_roi(frame, roi)
        lab = rgb_to_lab(mean_rgb)
        if reference_lab is None:
            reference_lab = lab.copy()
        row = {
            "target_s": target_time,
            "frame_s": actual_index / fps,
            "frame_index": actual_index,
            "r": float(mean_rgb[0]), "g": float(mean_rgb[1]),
            "b": float(mean_rgb[2]),
            "L": float(lab[0]), "a": float(lab[1]), "lab_b": float(lab[2]),
            "delta_e": delta_e_cie76(lab, reference_lab),
        }
        rows.append(row)
        print(f"{target_time:g} s (frame {actual_index}): DeltaE={row['delta_e']:.3f}")
        last_frame_index = target_index
        sample_number += 1
    if not rows:
        raise RuntimeError("No frames could be sampled.")
    return rows


def write_excel(path, rows, video_path, roi, interval_s, fps, frame_count):
    book = Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet.append(["Target time (s)", "Frame time (s)", "Frame index (0-based)",
                  "Mean R", "Mean G", "Mean B", "L*", "a*", "b*", "DeltaE (CIE76)"])
    for row in rows:
        sheet.append([row[k] for k in ("target_s", "frame_s", "frame_index",
                                      "r", "g", "b", "L", "a", "lab_b", "delta_e")])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for record in sheet.iter_rows(min_row=2):
        for i, cell in enumerate(record, start=1):
            if i != 3:
                cell.number_format = "0.000" if i > 3 else "0.000000"
    widths = [18, 18, 24, 12, 12, 12, 12, 12, 12, 19]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    settings = book.create_sheet("Settings")
    settings.append(["Parameter", "Value"])
    for cell in settings[1]:
        cell.font = Font(bold=True)
    for parameter, value in (
        ("Video file", str(video_path)),
        ("Reference frame", "First sampled frame (0 s)"),
        ("Sampling interval (s)", interval_s),
        ("Video FPS", fps),
        ("Video frame count", frame_count),
        ("Color pipeline", "Mean ROI RGB -> CIELAB (sRGB, D65)"),
        ("DeltaE formula", "CIE76 (DeltaE*ab)"),
        ("ROI x (px)", roi[0]), ("ROI y (px)", roi[1]),
        ("ROI width (px)", roi[2]), ("ROI height (px)", roi[3]),
    ):
        settings.append([parameter, value])
    settings.column_dimensions["A"].width = 29
    settings.column_dimensions["B"].width = 80
    book.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("video", nargs="?", help="Path to a video file; prompted if omitted")
    parser.add_argument("--interval", type=float, default=1.0,
                        help="Seconds between samples (default: 1)")
    parser.add_argument("--output", type=Path,
                        help="Excel output path (default: alongside video)")
    args = parser.parse_args(argv)
    raw = args.video or input("Enter the video file path: ").strip().strip('"')
    video_path = Path(raw).expanduser().resolve()
    if not video_path.is_file():
        parser.error(f"Video file does not exist: {video_path}")
    if video_path.suffix.lower() not in VIDEO_EXTENSIONS:
        print("Note: uncommon video extension; attempting to decode it anyway.")
    if not math.isfinite(args.interval) or args.interval <= 0:
        parser.error("--interval must be a positive finite number")
    output_path = args.output or video_path.with_name(video_path.stem + "_DeltaE_CIE76.xlsx")
    output_path = output_path.expanduser().resolve()
    cap = cv2.VideoCapture(str(video_path))
    if not cap.isOpened():
        raise RuntimeError(f"Cannot open video: {video_path}. Check the file and its codec.")
    try:
        fps, frame_count = video_info(cap)
        success, first_frame = cap.read()
        if not success:
            raise RuntimeError("Cannot decode the first frame of the video.")
        roi = RoiPlayer(cap, first_frame, fps, frame_count).select()
        print(f"Selected ROI: x={roi[0]}, y={roi[1]}, width={roi[2]}, height={roi[3]}")
        rows = analyze_video(cap, roi, args.interval)
    finally:
        cap.release()
    write_excel(output_path, rows, video_path, roi, args.interval, fps, frame_count)
    print(f"Saved {len(rows)} records to: {output_path}")


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError, RuntimeError, cv2.error) as exc:
        raise SystemExit(f"Error: {exc}") from exc
